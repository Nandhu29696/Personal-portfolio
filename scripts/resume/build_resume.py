"""Builds the resume (PDF + Word) and a cover letter template from content.json.

Outputs:
  public/Nandhakumar_M_Resume.pdf      linked from the portfolio
  career/Nandhakumar_M_Resume.docx     for employers or job boards that ask for Word
  career/Cover_Letter_Template.docx    fill in the highlighted [placeholders] per application

Follows the Job Bank (Government of Canada) resume guidance:
https://www.jobbank.gc.ca/findajob/resources/write-good-resume
  - two pages at most, contact details at the top
  - accomplishments with action verbs and numbers, 5-7 bullets per section at most
  - no photo, age, marital status, SIN, references, hobbies or personal pronouns
  - plain language and a simple single-column layout that ATS software can read

Run:  pip install reportlab python-docx
      python scripts/resume/build_resume.py
"""

import json
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[2]
C = json.loads((Path(__file__).with_name("content.json")).read_text(encoding="utf-8"))

PDF_OUT = ROOT / "public" / "Nandhakumar_M_Resume.pdf"
DOCX_OUT = ROOT / "career" / "Nandhakumar_M_Resume.docx"
LETTER_OUT = ROOT / "career" / "Cover_Letter_Template.docx"

BOLD = re.compile(r"\*\*(.+?)\*\*")
PLACEHOLDER = re.compile(r"(\[[^\]]+\])")


def html(text):
    """content.json text -> ReportLab markup (escaped, **bold** -> <b>)."""
    return BOLD.sub(r"<b>\1</b>", escape(text, quote=False))


# ================================================================ PDF (ReportLab)

INK, MUTED, ACCENT, RULE = HexColor("#18181b"), HexColor("#52525b"), HexColor("#a51c14"), HexColor("#d4d4d8")

base = ParagraphStyle("base", fontName="Helvetica", fontSize=9.6, leading=13, textColor=INK)
S = {
    "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=20, leading=24),
    "headline": ParagraphStyle("headline", parent=base, fontSize=11.5, leading=15, textColor=ACCENT),
    "contact": ParagraphStyle("contact", parent=base, fontSize=9, leading=12.5, textColor=MUTED),
    "section": ParagraphStyle("section", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=13,
                              textColor=ACCENT, spaceBefore=10, spaceAfter=2),
    "body": base,
    "role": ParagraphStyle("role", parent=base, fontName="Helvetica-Bold", fontSize=10.2),
    "meta": ParagraphStyle("meta", parent=base, textColor=MUTED),
    "right": ParagraphStyle("right", parent=base, textColor=MUTED, alignment=2),
    "bullet": ParagraphStyle("bullet", parent=base, alignment=TA_LEFT, leftIndent=11, bulletIndent=1,
                             spaceBefore=1.5),
}

# Frame width minus the 6pt padding reportlab puts on each side of the frame.
PAGE_WIDTH = LETTER[0] - 1.3 * inch - 12


def pdf_section(title, note=""):
    note = f' <font name="Helvetica" color="#52525b" size="9">{note}</font>' if note else ""
    return [Paragraph(title.upper() + note, S["section"]),
            HRFlowable(width="100%", thickness=0.6, color=RULE, spaceBefore=1, spaceAfter=5)]


def pdf_two_col(left, right, left_style="role"):
    t = Table([[Paragraph(left, S[left_style]), Paragraph(right, S["right"])]],
              colWidths=[PAGE_WIDTH * 0.72, PAGE_WIDTH * 0.28], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return t


def build_pdf():
    c = C["contact"]
    link = lambda text, href: f'<link href="{href}" color="#52525b">{html(text)}</link>'
    story = [
        Paragraph(html(c["name"]), S["name"]),
        Paragraph(html(c["headline"]), S["headline"]),
        Spacer(1, 3),
        Paragraph(html(c["location"]), S["contact"]),
        Paragraph(f'{html(c["phone"])} &middot; {link(c["email"], "mailto:" + c["email"])}', S["contact"]),
        Paragraph(" &middot; ".join(link(l["text"], l["href"]) for l in c["links"]), S["contact"]),
    ]

    story += pdf_section("Professional summary")
    story.append(Paragraph(html(C["summary"]), S["body"]))

    story += pdf_section("Technical skills")
    for label, value in C["skills"]:
        story.append(Paragraph(f"<b>{html(label)}:</b> {html(value)}", S["body"]))

    story += pdf_section("Professional experience")
    for job in C["experience"]:
        company = job["company"] + (f', {job["location"]}' if job["location"] else "")
        block = [pdf_two_col(html(job["role"]), html(job["dates"])), Paragraph(html(company), S["meta"])]
        block += [Paragraph(html(b), S["bullet"], bulletText="â€¢") for b in job["bullets"]]
        block.append(Spacer(1, 6))
        story.append(KeepTogether(block))

    story += pdf_section("Selected projects", "github.com/Nandhu29696")
    for name, stack, desc in C["projects"]:
        story.append(KeepTogether([
            Paragraph(f'<b>{html(name)}</b> <font color="#52525b">| {html(stack)}</font>', S["body"]),
            Paragraph(html(desc), S["body"]),
            Spacer(1, 4),
        ]))

    story += pdf_section("Certifications & training")
    for name, issuer, validity in C["certifications"]:
        story.append(pdf_two_col(f"<b>{html(name)}</b>, {html(issuer)}", html(validity), left_style="body"))

    story += pdf_section("Education")
    for degree, school, year in C["education"]:
        story.append(KeepTogether([pdf_two_col(html(degree), html(year)), Paragraph(html(school), S["meta"]),
                                   Spacer(1, 4)]))

    story += pdf_section("Awards")
    for name, org, why in C["awards"]:
        story.append(Paragraph(f"<b>{html(name)}</b>, {html(org)}: {html(why)}", S["body"]))

    SimpleDocTemplate(
        str(PDF_OUT), pagesize=LETTER,
        leftMargin=0.65 * inch, rightMargin=0.65 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch,
        title="Nandhakumar M - Resume", author="Nandhakumar M", subject="Full Stack AI Engineer resume",
    ).build(story)


# ================================================================ Word (python-docx)

D_INK, D_MUTED, D_ACCENT = RGBColor(0x18, 0x18, 0x1B), RGBColor(0x52, 0x52, 0x5B), RGBColor(0xA5, 0x1C, 0x14)
RIGHT_TAB = Inches(7.2)  # content width on Letter with 0.65" margins


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.65)
    section.top_margin = section.bottom_margin = Inches(0.55)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(10.5), D_INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08
    doc.core_properties.author = "Nandhakumar M"
    return doc


def add_runs(paragraph, text, color=None, size=None, highlight_placeholders=False):
    """Adds text to a paragraph, honouring **bold** and optionally highlighting [placeholders]."""
    for i, chunk in enumerate(BOLD.split(text)):
        parts = PLACEHOLDER.split(chunk) if highlight_placeholders else [chunk]
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run(part)
            run.bold = i % 2 == 1
            if color is not None:
                run.font.color.rgb = color
            if size is not None:
                run.font.size = Pt(size)
            if highlight_placeholders and PLACEHOLDER.fullmatch(part):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def add_hyperlink(paragraph, text, url, color=D_MUTED, size=9.5):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), str(color))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.extend([col, sz])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def bottom_border(paragraph, color="D4D4D8"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", color)):
        bottom.set(qn(f"w:{key}"), val)
    borders.append(bottom)
    ppr.append(borders)


def contact_header(doc, compact=False):
    c = C["contact"]
    p = doc.add_paragraph()
    add_runs(p, c["name"], size=20 if not compact else 18).runs[0].bold = True
    add_runs(doc.add_paragraph(), c["headline"], color=D_ACCENT, size=12)
    add_runs(doc.add_paragraph(), c["location"], color=D_MUTED, size=9.5)
    p = doc.add_paragraph()
    add_runs(p, f'{c["phone"]} Â· ', color=D_MUTED, size=9.5)
    add_hyperlink(p, c["email"], f'mailto:{c["email"]}')
    p = doc.add_paragraph()
    for i, l in enumerate(c["links"]):
        if i:
            add_runs(p, " Â· ", color=D_MUTED, size=9.5)
        add_hyperlink(p, l["text"], l["href"])
    return p


def docx_section(doc, title, note=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    add_runs(p, title.upper(), color=D_ACCENT, size=11).runs[0].bold = True
    if note:
        add_runs(p, f"  {note}", color=D_MUTED, size=9.5)
    bottom_border(p)


def left_right(doc, left, right, bold_left=True):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left)
    run.bold = bold_left
    tail = p.add_run(f"\t{right}")
    tail.font.color.rgb = D_MUTED
    return p


def build_resume_docx():
    doc = new_document()
    contact_header(doc)

    docx_section(doc, "Professional summary")
    add_runs(doc.add_paragraph(), C["summary"])

    docx_section(doc, "Technical skills")
    for label, value in C["skills"]:
        add_runs(doc.add_paragraph(), f"**{label}:** {value}")

    docx_section(doc, "Professional experience")
    for job in C["experience"]:
        p = left_right(doc, job["role"], job["dates"])
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        company = job["company"] + (f', {job["location"]}' if job["location"] else "")
        meta = add_runs(doc.add_paragraph(), company, color=D_MUTED)
        meta.paragraph_format.keep_with_next = bool(job["bullets"])
        for b in job["bullets"]:
            add_runs(doc.add_paragraph(style="List Bullet"), b)

    docx_section(doc, "Selected projects", "github.com/Nandhu29696")
    for name, stack, desc in C["projects"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.keep_with_next = True
        add_runs(p, f"**{name}**")
        add_runs(p, f" | {stack}", color=D_MUTED)
        add_runs(doc.add_paragraph(), desc)

    docx_section(doc, "Certifications & training")
    for name, issuer, validity in C["certifications"]:
        left_right(doc, f"{name}, {issuer}", validity)

    docx_section(doc, "Education")
    for degree, school, year in C["education"]:
        p = left_right(doc, degree, year)
        p.paragraph_format.keep_with_next = True
        add_runs(doc.add_paragraph(), school, color=D_MUTED).paragraph_format.space_after = Pt(4)

    docx_section(doc, "Awards")
    for name, org, why in C["awards"]:
        add_runs(doc.add_paragraph(), f"**{name}**, {org}: {why}")

    doc.core_properties.title = "Nandhakumar M - Resume"
    doc.save(DOCX_OUT)


def build_cover_letter():
    doc = new_document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15

    last = contact_header(doc, compact=True)
    last.paragraph_format.space_after = Pt(18)
    bottom_border(last)

    L = C["cover_letter"]
    for line in L["recipient"]:
        add_runs(doc.add_paragraph(), line, highlight_placeholders=True)
    doc.add_paragraph()

    subject = add_runs(doc.add_paragraph(), L["subject"], highlight_placeholders=True)
    for run in subject.runs:
        run.bold = True
    subject.paragraph_format.space_after = Pt(12)

    add_runs(doc.add_paragraph(), L["greeting"], highlight_placeholders=True).paragraph_format.space_after = Pt(10)
    for text in L["paragraphs"]:
        add_runs(doc.add_paragraph(), text, highlight_placeholders=True).paragraph_format.space_after = Pt(10)

    doc.add_paragraph(L["closing"]).paragraph_format.space_after = Pt(24)
    doc.add_paragraph(L["signature"])

    doc.core_properties.title = "Cover letter template - Nandhakumar M"
    doc.save(LETTER_OUT)


if __name__ == "__main__":
    DOCX_OUT.parent.mkdir(exist_ok=True)
    build_pdf()
    build_resume_docx()
    build_cover_letter()
    for path in (PDF_OUT, DOCX_OUT, LETTER_OUT):
        print(f"Wrote {path.relative_to(ROOT)}")
